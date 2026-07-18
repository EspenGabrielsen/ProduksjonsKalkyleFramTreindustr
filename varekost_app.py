import marimo

__generated_with = "0.23.14"
app = marimo.App(width="medium")


@app.cell
def _():
    import marimo as mo
    import sys
    import io
    import os
    import re
    import tempfile
    import shutil
    from datetime import datetime, date
    from pathlib import Path
    from typing import Optional

    import pandas as pd

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    elif hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

    from kostberegning import (
        ExcelData,
        CostCalculator,
        SimulationEngine,
        SimulationOverride,
        SimulationComparison,
        ProductCostResult,
        MaterialCostDetail,
        OperationCostDetail,
        ByProductDetail,
        export_product_costs_to_json,
    )

    # Importer PDF-rapport-generator
    from generer_simuleringsrapport import generer_rapport, registrer_fonter, _hent_logo

    return (
        CostCalculator,
        ExcelData,
        SimulationEngine,
        SimulationOverride,
        generer_rapport,
        mo,
        os,
        pd,
        registrer_fonter,
        tempfile,
        _hent_logo,
    )


@app.cell
def _(mo):
    # Global CSS for Fram Treindustri-profil
    mo.Html("""
    <style>
        body { 
            /* En veldig lys, behagelig grågrønn/varm hvit bakgrunn */
            background-color: #F3F5F2; 
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            color: #2C3E2B; /* Mørk skoggrønn i stedet for svart tekst for et mykere uttrykk */
        }
    
        .marimo-app { 
            max-width: 1200px; 
            margin: 40px auto; 
            padding: 0 24px;
        }
    
        .fti-card {
            /* Hvite kort med en ørliten nyanse av lysegrønt, som gir fin dybde mot bakgrunnen */
            background: #F9FBF8; 
            border-radius: 12px; 
            padding: 24px;
            margin-bottom: 20px;
        
            /* En myk skygge som gjør at kortene "svever" lett */
            box-shadow: 0 4px 12px rgba(27, 89, 43, 0.04);
        
            /* Den friske, lysegrønne FramTre-aksenten på venstre side */
            border-left: 5px solid #48BB78; 
        }
    
        .fti-card h2, .fti-card h3 {
            /* Dyp skogsgrønn på overskrifter for god lesbarhet og kontrast */
            color: #14532D; 
            margin-top: 0;
            margin-bottom: 12px;
            font-weight: 600;
        }
    
        /* Highlight-farger tilpasset den nye lysegrønne profilen: */
        .fti-highlight-green { 
            color: #2F855A; 
            background-color: #E6FFFA; /* Subtil grønn merking bak teksten */
            padding: 2px 6px;
            border-radius: 4px;
            font-weight: 600; 
        }
    
        .fti-highlight-red { 
            color: #C53030; 
            background-color: #FFF5F5; /* Subtil rød merking bak teksten */
            padding: 2px 6px;
            border-radius: 4px;
            font-weight: 600; 
        }
    </style>
    """)
    return


@app.cell
def _(mo):
    mo.hstack([
        mo.image(
            src="https://framtreindustri.no/wp-content/uploads/2025/08/logo-liggende-2048x512.png",
            alt="Fram Treindustri",
            width=250,
        ),
    mo.Html(
            '<div style="text-align: right; line-height: 1.2;">'
            '<div style="font-size: 1.8em; font-weight: 700; color: #1B3A5C;">Produksjonskost Simulator</div>'
            '<div style="font-size: 0.95em; color: #2C5F8A;">Last opp Excel-modell, juster parametere — se kostnader i sanntid</div>'
            '</div>'
        ),
    ], justify="space-between", align="center")
    return


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>📂 Last data</h2></div>')
    return


@app.cell
def _(mo):
    file_upload = mo.ui.file(
        label="Velg Excel-fil (Produksjonsmodell)",
        filetypes=[".xlsx"],
        multiple=False,
    )
    file_upload
    return (file_upload,)


@app.cell
def _():
    persistent_overrides = {
        "item_costs": {},
        "bom_scrap": {},
        "bom_co_product": {},
        "work_centers": {},
        "routing": {},
    }
    return (persistent_overrides,)


@app.cell
def _(
    CostCalculator,
    ExcelData,
    SimulationEngine,
    file_upload,
    mo,
    os,
    tempfile,
):
    data = None
    engine = None
    baseline = None

    if file_upload.value:
        try:
            _upload = file_upload.value[0]
            _fname = _upload.name
            _fcontent = _upload.contents

            _temp_dir = tempfile.mkdtemp()
            _temp_path = os.path.join(_temp_dir, _fname)
            with open(_temp_path, "wb") as _f:
                _f.write(_fcontent)

            data = ExcelData(_temp_path)
            engine = SimulationEngine(data)
            _calculator = CostCalculator(data)
            baseline = _calculator.calculate_all()

            _n_fg = len([p for p in data.products if p.item_type in ('Finished Good', 'Semi Finished')])
            _n_rm = len([p for p in data.products if p.item_type == 'Raw Material'])
            _n_bp = len([p for p in data.products if p.item_type == 'By Product'])

            mo.output.replace(
                mo.md(
                    f"""
                    ### ✅ Data lastet!

                    | Data | Antall |
                    |---|---|
                    | Produkter | {len(data.products)} |
                    | Ferdigvarer/Halvfabrikata | {_n_fg} |
                    | Råvarer | {_n_rm} |
                    | Biprodukter | {_n_bp} |
                    | Arbeidssentre | {len(data.work_centers)} |
                    | BOM-linjer | {len(data.bom_lines)} |
                    | Routing-linjer | {len(data.routing_lines)} |
                    | Biprodukt-regler | {len(data.byproduct_rules)} |
                    | Scenarioer | {len(data.scenarios)} |
                    """
                )
            )

        except Exception as _e:
            mo.output.replace(mo.md(f"### ❌ Feil ved lasting av fil: {_e}"))
            import traceback as _traceback
            _traceback.print_exc()
    return baseline, data


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>🔍 Vare filter</h2></div>')
    return


@app.cell
def _(file_upload, mo):
    if file_upload.value:
        vareFilter = mo.ui.text(label="🔍 Vare filter (søk på varenr eller beskrivelse)")
    else: 
        vareFilter = mo.ui.text(label="🔍 Vare filter (Last inn Excel først)", disabled=True)
    vareFilter
    return (vareFilter,)


@app.cell
def _(data, vareFilter):
    # Returner filtrerte datasett basert på vareFilter
    # Kaskade: filtrer produkter → filtrer BOM/routing → filtrer WC/lokasjoner
    filtered_products = []
    filtered_bom_lines = []
    filtered_routing_lines = []
    filtered_byproduct_rules = []
    filtered_item_costs = []
    filtered_scenarios = []
    filtered_work_centers = []
    filtered_locations = []
    filtered_operations = []
    filtered_capacity_days = []
    filter_active = False

    if data:
        _filter_text = vareFilter.value.strip().lower() if vareFilter.value else ""

        if _filter_text:
            filter_active = True

            # 1. Finn FG/Semi Finished som matcher filteret (item_no eller description LIKE '%filter%')
            _fg_set = set()
            for _p in data.products:
                if _p.item_type in ('Finished Good', 'Semi Finished'):
                    if _filter_text in _p.item_no.lower() or _filter_text in _p.description.lower():
                        _fg_set.add(_p.item_no)

            # 2. Filtrer BOM: kun linjer der parent er i _fg_set
            #    Samle opp komponentene (råvarer) som trengs
            _component_set = set()
            for _bl in data.bom_lines:
                if _bl.parent_item_no in _fg_set:
                    filtered_bom_lines.append(_bl)
                    _component_set.add(_bl.component_item_no)

            # 3. Bygg det komplette produktsettet: FG + komponenter (råvarer) + biprodukter
            _all_product_set = set(_fg_set)

            # Legg til råvarer som er komponenter i BOM for filtrerte FG
            for _p in data.products:
                if _p.item_type == 'Raw Material' and _p.item_no in _component_set:
                    _all_product_set.add(_p.item_no)

            # Legg til biprodukter som oppstår fra filtrerte FG
            _bp_set = set()
            for _br in data.byproduct_rules:
                if _br.parent_item_no in _fg_set:
                    _bp_set.add(_br.by_product_item_no)
            for _p in data.products:
                if _p.item_type == 'By Product' and _p.item_no in _bp_set:
                    _all_product_set.add(_p.item_no)

            # 4. Bygg filtered_products (kun produkter i _all_product_set)
            for _p in data.products:
                if _p.item_no in _all_product_set:
                    filtered_products.append(_p)

            # 5. Filtrer byproduct rules (parent i _fg_set)
            for _br in data.byproduct_rules:
                if _br.parent_item_no in _fg_set:
                    filtered_byproduct_rules.append(_br)

            # 6. Filtrer routing (kun for FG/Semi Finished i _fg_set)
            _routing_wc_set = set()
            for _rl in data.routing_lines:
                if _rl.item_no in _fg_set:
                    filtered_routing_lines.append(_rl)
                    _routing_wc_set.add(_rl.work_center_code)

            # 7. Filtrer item costs (item_no i _all_product_set)
            for _ic in data.item_costs:
                if _ic.item_no in _all_product_set:
                    filtered_item_costs.append(_ic)

            # 8. Filtrer scenarios (product i _fg_set)
            for _sc in data.scenarios:
                if _sc.product in _fg_set:
                    filtered_scenarios.append(_sc)

            # 9. Filtrer work centers (de som finnes i filtered routing)
            for _wc in data.work_centers:
                if _wc.code in _routing_wc_set:
                    filtered_work_centers.append(_wc)

            # 10. Filtrer locations (de som finnes i filtered work centers)
            _loc_set = {_wc.location_code for _wc in filtered_work_centers}
            for _loc in data.locations:
                if _loc.code in _loc_set:
                    filtered_locations.append(_loc)

            # 11. Filtrer operations (de som finnes i filtered routing)
            _op_set = {_rl.operation_code for _rl in filtered_routing_lines}
            for _op in data.operations:
                if _op.code in _op_set:
                    filtered_operations.append(_op)

            # 12. Filtrer capacity days (for filtered work centers)
            for _cd in data.capacity_days:
                if _cd.work_center in _routing_wc_set:
                    filtered_capacity_days.append(_cd)
        else:
            # Ingen filter = alle data
            filtered_products = list(data.products)
            filtered_bom_lines = list(data.bom_lines)
            filtered_routing_lines = list(data.routing_lines)
            filtered_byproduct_rules = list(data.byproduct_rules)
            filtered_item_costs = list(data.item_costs)
            filtered_scenarios = list(data.scenarios)
            filtered_work_centers = list(data.work_centers)
            filtered_locations = list(data.locations)
            filtered_operations = list(data.operations)
            filtered_capacity_days = list(data.capacity_days)
    return (
        filter_active,
        filtered_bom_lines,
        filtered_byproduct_rules,
        filtered_capacity_days,
        filtered_item_costs,
        filtered_locations,
        filtered_operations,
        filtered_products,
        filtered_routing_lines,
        filtered_scenarios,
        filtered_work_centers,
    )


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>📊 Baseline kostnader</h2><p style="color:#2C5F8A;">Original kalkylestruktur (før simulering)</p></div>')
    return


@app.cell
def _(baseline, mo, pd):
    if baseline:
        _rows = []
        for _r in baseline:
            _rows.append({
                "Lokasjon": _r.location_code,
                "Produkt": _r.product_no,
                "Beskrivelse": _r.product_desc,
                "Materialkost": round(_r.material_cost, 2),
                "Operasjonskost": round(_r.operation_cost, 2),
                "Setupkost": round(_r.setup_cost, 2),
                "Brutto kost": round(_r.gross_production_cost, 2),
                "Biproduktverdi": round(_r.by_product_value, 2),
                "Netto kost": round(_r.net_production_cost, 2),
            })
        _df = pd.DataFrame(_rows)
        mo.output.replace(mo.ui.table(_df, selection=None))
    return


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>📋 Originale data fra modellen</h2></div>')
    return


@app.cell
def _(mo):
    mo.md("""
    **📋 Dataark - oversikt**

    Modellen består av **10 ark** i Excel. Hvert ark har en spesifikk rolle:

    | # | Arknavn | Innhold | Nøkkelkolonner |
    |---|---------|---------|----------------|
    | 1 | **Product Master** | Vareregister | Item No, Description, Item Type |
    | 2 | **Locations** | Fabrikker og lagre | Location Code, Location Name |
    | 3 | **Work Centers** | Arbeidssentre med kostsatser | Work Center Code, Labor/Machine/Overhead Cost |
    | 4 | **Operation Master** | Standardoperasjoner | Operation Code, Description |
    | 5 | **Item Costs** | Kostpriser per vare | Item No, Unit Cost, Currency |
    | 6 | **BOM** | Stykkliste (hva består produktet av) | Parent Item, Component, Quantity Per |
    | 7 | **Routing** | Produksjonsflyt (operasjoner, tider) | Item No, Operation No, Setup/Run Time |
    | 8 | **By Product Rules** | Biprodukter og verdsetting | Parent Item, By Product, Market Value |
    | 9 | **Capacity Calendar** | Kapasitetskalender per arbeidssenter | Work Center, Date, Available Hours |
    | 10 | **Production Scenario** | Produksjonsscenarioer | Scenario Name, Product, Planned Quantity |
    """)
    return


@app.cell
def _(
    data,
    filter_active,
    filtered_bom_lines,
    filtered_byproduct_rules,
    filtered_capacity_days,
    filtered_item_costs,
    filtered_locations,
    filtered_operations,
    filtered_products,
    filtered_routing_lines,
    filtered_scenarios,
    filtered_work_centers,
    mo,
    pd,
):
    if data:
        _outputs = []

        # Vis filter-status
        if filter_active:
            _outputs.append(mo.md(f"**🔍 Filter aktiv:** {len(filtered_products)} produkter matcher"))

        # ── Product Master ──────────────────────────────────────────
        _outputs.append(mo.md("""
        **📦 Product Master (Vareregister)**

        Register over alle varer i virksomheten. Hovedkatalogen over råvarer, halvfabrikata, ferdigvarer, biprodukter og handelsvarer.

        | Kolonne | Type | Beskrivelse | Eksempel |
        |---------|------|-------------|----------|
        | **Item No** | Tekst | Unik identifikator for varen | `RM001`, `FG001`, `BP001` |
        | **Description** | Tekst | Beskrivende navn på varen | `Gran 50x200 US/V`, `Utvendig Panel 21x148` |
        | **Item Type** | Tekst | Type vare: `Raw Material`, `Semi Finished`, `Finished Good`, `By Product`, `Trading Item` | `Raw Material` |
        | **Product Group** | Tekst | Gruppering av varer | `Skrulast`, `Panel`, `Kledning`, `Spon` |
        | **Base Unit of Measure** | Tekst | Standard måleenhet | `LM`, `M3`, `KG`, `LTR` |
        | **Active** | Ja/Nei | Angir om varen er aktiv | `Ja` |
        """))

        # Produkter (filtrert)
        _prod_rows = []
        for _p in filtered_products:
            _prod_rows.append({
                "Varenr": _p.item_no,
                "Beskrivelse": _p.description,
                "Type": _p.item_type,
                "Gruppe": _p.product_group,
                "Enhet": _p.base_uom,
                "Aktiv": "Ja" if _p.active else "Nei",
            })
        _outputs.append(mo.ui.table(pd.DataFrame(_prod_rows), selection=None))

        # ── Locations ───────────────────────────────────────────────
        if filtered_locations:
            _outputs.append(mo.md("""
            **🏭 Locations (Lokasjoner)**

            Register over fabrikker og lagre. Hvert arbeidssenter er knyttet til en lokasjon.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Location Code** | Tekst | Unik kode for lokasjonen | `KOD` |
            | **Location Name** | Tekst | Navn på lokasjonen | `Kodal Fabrikk` |
            | **Location Type** | Tekst | Type lokasjon: `Factory`, `Warehouse`, `Distribution Center`, `Sales Office` | `Factory` |
            | **Active** | Ja/Nei | Angir om lokasjonen er aktiv | `Ja` |
            """))
            _loc_rows = []
            for _loc in filtered_locations:
                _loc_rows.append({
                    "Kode": _loc.code,
                    "Navn": _loc.name,
                    "Type": _loc.location_type,
                    "Aktiv": "Ja" if _loc.active else "Nei",
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_loc_rows), selection=None))

        # ── Work Centers ────────────────────────────────────────────
        if filtered_work_centers:
            _outputs.append(mo.md("""
            **🏭 Work Centers (Arbeidssentre)**

            Register over produksjonsressurser - maskiner og arbeidsplasser. Hvert arbeidssenter har timekostnader som brukes til å beregne operasjonskost.

            **Timekostnad:** `Total kost per time = Lønn + Maskin + Overhead`

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Work Center Code** | Tekst | Unik identifikator | `HOVEDHOVEL` |
            | **Description** | Tekst | Beskrivende navn | `Hovedhovel` |
            | **Location Code** | Tekst | Fabrikken det tilhører | `KOD` |
            | **Labor Cost per Hour** | Desimal | Arbeidskostnad per time (lønn, arbeidsgiveravgift, pensjon, feriepenger) | `550` |
            | **Machine Cost per Hour** | Desimal | Maskinkostnad per time (avskrivninger, service, leasing, vedlikehold, energi) | `900` |
            | **Overhead Cost per Hour** | Desimal | Indirekte produksjonskostnader (produksjonsledelse, kvalitet, vedlikeholdsadm., intern logistikk) | `150` |
            | **Capacity Hours per Day** | Desimal | Tilgjengelige timer per dag | `16` |
            | **Effective Capacity %** | Prosent | Hvor stor del av tiden som faktisk kan brukes til produksjon | `85` |
            | **Active** | Ja/Nei | Angir om arbeidssenteret er aktivt | `Ja` |
            """))
            _wc_rows = []
            for _wc in filtered_work_centers:
                _wc_rows.append({
                    "Arbeidssenter": _wc.code,
                    "Beskrivelse": f"{_wc.description} · {_wc.location_code}",
                    "Lønn/time": _wc.labor_cost_hour,
                    "Maskin/time": _wc.machine_cost_hour,
                    "Overhead/time": _wc.overhead_cost_hour,
                    "Totalt/time": _wc.total_cost_hour,
                    "Timer/dag": _wc.capacity_hours_day,
                    "Eff. %": _wc.effective_capacity_pct,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_wc_rows), selection=None))

        # ── Operation Master ────────────────────────────────────────
        if filtered_operations:
            _outputs.append(mo.md("""
            **⚙️ Operation Master (Operasjoner)**

            Standardisert liste over operasjoner som kan brukes i routing. Gir en felles "ordbok" for produksjonsprosesser.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Operation Code** | Tekst | Unik operasjonskode | `HOVLING`, `MALING`, `PACKING` |
            | **Description** | Tekst | Beskrivelse av operasjonen | `Høvling (oppdeling + høvling + profilering)` |
            | **Default Work Center** | Tekst | Anbefalt arbeidssenter | `HOVEDHOVEL` |
            | **Standard Unit** | Tekst | Måleenhet for produksjonstid | `Minutes`, `Hours` |
            | **Active** | Ja/Nei | Angir om operasjonen er aktiv | `Ja` |
            """))
            _op_rows = []
            for _op in filtered_operations:
                _op_rows.append({
                    "Kode": _op.code,
                    "Beskrivelse": _op.description,
                    "Standard arbeidssenter": _op.default_work_center,
                    "Enhet": _op.standard_unit,
                    "Aktiv": "Ja" if _op.active else "Nei",
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_op_rows), selection=None))

        # ── Item Costs ──────────────────────────────────────────────
        if filtered_item_costs:
            _outputs.append(mo.md("""
            **💰 Item Costs (Kostpriser)**

            Samlet register over kostpriser for alle varer. For råvarer er dette innkjøpspris. For ferdigvarer settes prisen til 0 (beregnes automatisk). For biprodukter er dette markedsverdi.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Item No** | Tekst | Referanse til varen (fra Product Master) | `RM001` |
            | **Cost Type** | Tekst | Type kostpris: `Standard Cost`, `Last Direct Cost`, `Forecast Cost`, `Budget Cost` | `Standard Cost` |
            | **Unit Cost** | Desimal | Kostpris per enhet | `3390.00` |
            | **Currency** | Tekst | Valuta | `NOK`, `EUR` |
            | **Effective Date** | Dato | Dato kostprisen gjelder fra | `2026-01-01` |
            """))
            _ic_rows = []
            for _ic in filtered_item_costs:
                _ic_rows.append({
                    "Varenr": _ic.item_no,
                    "Kosttype": _ic.cost_type,
                    "Enhetskost": _ic.unit_cost,
                    "Valuta": _ic.currency,
                    "Gyldig fra": _ic.effective_date,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_ic_rows), selection=None))

        # ── BOM ─────────────────────────────────────────────────────
        if filtered_bom_lines:
            _outputs.append(mo.md("""
            **🔗 BOM (Stykkliste)**

            Beskriver hvilke komponenter som inngår i et produkt. En BOM-linje sier: "For å lage X trenger du Y".

            **Quantity Per** = antall output-enheter per input-enhet.  
            **Forbruk per output** = 1 / Quantity Per.  
            **Materialkost per enhet** = (Unit Cost / Quantity Per) × (1 + Scrap% / 100).

            **Co-Prod %** = andel av produksjonen som blir samprodukt (co-product).  
            Co-produkter (f.eks. B-vare) får sin egen kalkyle med allokert material- og operasjonskost.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Parent Item No** | Tekst | Produktet som produseres | `FG001` |
            | **Component Item No** | Tekst | Komponenten som forbrukes | `RM002` |
            | **Quantity Per** | Desimal | Antall output-enheter per input-enhet | `303.04` |
            | **Unit of Measure** | Tekst | Måleenhet for forholdet | `LM` |
            | **Scrap %** | Prosent | Forventet materialsvinn | `5.0` |
            | **Co-Prod %** | Prosent | Andel som blir samprodukt (co-product) | `6.0` |
            | **Co-Prod Item No** | Tekst | Varenummer for samproduktet | `JD16073-B` |
            | **Valid From** | Dato | Gyldig fra dato | `2026-01-01` |
            | **Valid To** | Dato | Gyldig til dato | `2026-12-31` |
            """))
            _bom_rows = []
            for _bl in filtered_bom_lines:
                _bom_rows.append({
                    "Foreldreprodukt": _bl.parent_item_no,
                    "Komponent": _bl.component_item_no,
                    "Qty per": _bl.quantity_per,
                    "Enhet": _bl.uom,
                    "Svinn %": _bl.scrap_pct,
                    "Co-Prod %": _bl.co_product_pct,
                    "Co-Prod Vare": _bl.co_product_item_no,
                    "Gyldig fra": _bl.valid_from,
                    "Gyldig til": _bl.valid_to,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_bom_rows), selection=None))

        # ── Routing ─────────────────────────────────────────────────
        if filtered_routing_lines:
            _outputs.append(mo.md("""
            **📋 Routing (Produksjonsflyt)**

            Beskriver produksjonsprosessen - hvilke operasjoner som utføres, i hvilken rekkefølge, på hvilket arbeidssenter, og hvor lang tid hver operasjon tar.

            **Operasjonskost per enhet** = (Run Time Minutes / 60) × Timekost.  
            **Setupkost per enhet** = ((Setup Time Minutes / 60) × Timekost) / Batch Size.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Item No** | Tekst | Produkt som produseres | `FG001` |
            | **Operation No** | Heltall | Sekvensnummer (stigende rekkefølge) | `10`, `20`, `30` |
            | **Operation Code** | Tekst | Hvilken operasjon som utføres | `HOVLING`, `MALING` |
            | **Work Center Code** | Tekst | Arbeidssenter som utfører operasjonen | `HOVEDHOVEL` |
            | **Setup Time Minutes** | Desimal | Tid til klargjøring (omstilling, knivbytte, innkjøring, kontrollmåling) | `15.0` |
            | **Run Time Minutes** | Desimal | Produksjonstid per enhet | `0.03` |
            | **Batch Size** | Desimal | Normal ordrestørrelse (brukes til å fordele setupkostnad) | `50000` |
            | **Valid From** | Dato | Gyldig fra dato | `2026-01-01` |
            | **Valid To** | Dato | Gyldig til dato | `2026-12-31` |
            """))
            _rt_rows = []
            for _rl in filtered_routing_lines:
                _op = next((o for o in filtered_operations if o.code == _rl.operation_code), None)
                _op_desc = _op.description if _op else _rl.operation_code
                _wc = next((w for w in filtered_work_centers if w.code == _rl.work_center_code), None)
                _loc = _wc.location_code if _wc else ""
                _rt_rows.append({
                    "Produkt": _rl.item_no,
                    "Op.nr": _rl.operation_no,
                    "Operasjon": _rl.operation_code,
                    "Beskrivelse": _op_desc,
                    "Arbeidssenter": _rl.work_center_code,
                    "Lokasjon": _loc,
                    "Setup (min)": _rl.setup_time_minutes,
                    "Run time (min)": _rl.run_time_minutes,
                    "Batch": _rl.batch_size,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_rt_rows), selection=None))

        # ── By Product Rules ────────────────────────────────────────
        if filtered_byproduct_rules:
            _outputs.append(mo.md("""
            **♻️ By Product Rules (Biproduktregler)**

            Beskriver hvordan biprodukter håndteres økonomisk. I trelastproduksjon oppstår det alltid biprodukter som høvelspon, flis og bark. Disse har en verdi som skal trekkes fra hovedproduktets kostnad.

            **Biproduktverdi per enhet** = Expected Quantity × Market Value.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Parent Item No** | Tekst | Produktet (ferdigvaren) som skaper biproduktet | `FG001` |
            | **By Product Item No** | Tekst | Biproduktet | `BP001` |
            | **Expected Quantity** | Desimal | Forventet mengde biprodukt per enhet hovedprodukt | `0.15` |
            | **Unit of Measure** | Tekst | Måleenhet for biproduktet | `KG` |
            | **Market Value** | Desimal | Forventet markedspris per enhet | `1.50` |
            | **Allocation Method** | Tekst | Hvordan verdien skal håndteres: `Reduce Main Product Cost`, `Separate Profit Center`, `Informational Only` | `Reduce Main Product Cost` |
            """))
            _bp_rows = []
            for _br in filtered_byproduct_rules:
                _bp_rows.append({
                    "Produkt": _br.parent_item_no,
                    "Biprodukt": _br.by_product_item_no,
                    "Forventet qty": _br.expected_quantity,
                    "Enhet": _br.uom,
                    "Markedsverdi": _br.market_value,
                    "Allokeringsmetode": _br.allocation_method,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_bp_rows), selection=None))

        # ── Capacity Calendar ───────────────────────────────────────
        if filtered_capacity_days:
            _outputs.append(mo.md("""
            **📅 Capacity Calendar (Kapasitetskalender)**

            Viser tilgjengelig kapasitet per arbeidssenter per dag. Brukes til å analysere flaskehalser og kapasitetsutnyttelse.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Work Center** | Tekst | Arbeidssenter | `HOVEDHOVEL` |
            | **Date** | Dato | Dato | `2026-01-05` |
            | **Available Hours** | Desimal | Tilgjengelige timer per dag | `16` |
            | **Planned Downtime** | Desimal | Planlagt nedetid (vedlikehold, stopp) | `0` |
            | **Available Production Hours** | Desimal | Faktisk produksjonstid = Available Hours - Planned Downtime | `16` |
            """))
            _cap_rows = []
            for _cd in filtered_capacity_days:
                _cap_rows.append({
                    "Arbeidssenter": _cd.work_center,
                    "Dato": _cd.date,
                    "Tilgj. timer": _cd.available_hours,
                    "Planlagt nedetid": _cd.planned_downtime,
                    "Prod. timer": _cd.available_production_hours,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_cap_rows), selection=None))

        # ── Production Scenario ─────────────────────────────────────
        if filtered_scenarios:
            _outputs.append(mo.md("""
            **🎯 Production Scenario (Produksjonsscenarioer)**

            Forhåndsdefinerte produksjonsscenarioer med planlagt kvantum per produkt. Brukes til å simulere produksjon og beregne totalt ressursbehov.

            | Kolonne | Type | Beskrivelse | Eksempel |
            |---------|------|-------------|----------|
            | **Scenario Name** | Tekst | Navn på scenario | `Normal Produksjon` |
            | **Product** | Tekst | Produkt som skal produseres | `FG001` |
            | **Planned Quantity** | Desimal | Planlagt antall enheter | `50000` |
            | **Start Date** | Dato | Startdato for produksjon | `2026-01-01` |
            | **End Date** | Dato | Sluttdato for produksjon | `2026-12-31` |
            """))
            _sc_rows = []
            for _sc in filtered_scenarios:
                _sc_rows.append({
                    "Scenario": _sc.scenario_name,
                    "Produkt": _sc.product,
                    "Planlagt kvantum": _sc.planned_quantity,
                    "Startdato": _sc.start_date,
                    "Sluttdato": _sc.end_date,
                })
            _outputs.append(mo.ui.table(pd.DataFrame(_sc_rows), selection=None))

        mo.output.replace(
            mo.accordion({
                "🔍 Vis detaljerte stamdata fra Excel-modellen (BOM, Routing, Arbeidssentre osv.)": mo.vstack(_outputs)
            })
        )
    return


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>🔧 Juster parametere</h2></div>')
    return


@app.cell(hide_code=True)
def _(mo):
    mo.md(r"""
    **🪵 Råvarer**
    """)
    return


@app.cell
def _(data, filtered_products, mo, pd, persistent_overrides):
    # Redigerbar tabell for råvarepriser
    rm_price_df = None
    if data:
        _rm_items = [p for p in filtered_products if p.item_type == 'Raw Material']
        _rows = []
        for _p in _rm_items:
            _cost = next((c.unit_cost for c in data.item_costs if c.item_no == _p.item_no), 0)
            _ny_cost = persistent_overrides["item_costs"].get(_p.item_no, _cost)
            _rows.append({
                "Varenr": _p.item_no,
                "Beskrivelse": f"{_p.description} · {_p.base_uom}",
                "Org. pris": _cost,
                "Ny pris": _ny_cost,
            })
        if _rows:
            _df = pd.DataFrame(_rows)
            rm_price_df = mo.ui.data_editor(_df, editable_columns=["Ny pris"])

    rm_price_df
    return (rm_price_df,)


@app.cell
def _(persistent_overrides, rm_price_df):
    if rm_price_df is not None and rm_price_df.value is not None:
        _df = rm_price_df.value
        for _, _row in _df.iterrows():
            _varenr = _row["Varenr"]
            _org = _row["Org. pris"]
            _ny = _row["Ny pris"]
            if abs(_ny - _org) > 0.001:
                persistent_overrides["item_costs"][_varenr] = _ny
            elif _varenr in persistent_overrides["item_costs"]:
                del persistent_overrides["item_costs"][_varenr]
    return


@app.cell
def _(mo):
    mo.md("""
    **🗑️ Svinn- og kapp-prosenter**
    """)
    return


@app.cell
def _(data, filtered_bom_lines, mo, pd, persistent_overrides):
    # Redigerbar tabell for svinn-prosenter og co-prod %
    bom_scrap_df = None
    if data:
        _rows = []
        for _bl in filtered_bom_lines:
            _key = (_bl.parent_item_no, _bl.component_item_no)
            _ny_svinn = persistent_overrides["bom_scrap"].get(_key, _bl.scrap_pct)
            _ny_co = persistent_overrides["bom_co_product"].get(_key, _bl.co_product_pct)
            # Slå opp beskrivelser for komponent og produkt
            _komp_desc = next((p.description for p in data.products if p.item_no == _bl.component_item_no), _bl.component_item_no)
            _prod_desc = next((p.description for p in data.products if p.item_no == _bl.parent_item_no), _bl.parent_item_no)
            _rows.append({
                "Komponent": f"{_bl.component_item_no} · {_komp_desc}",
                "Produkt": f"{_bl.parent_item_no} · {_prod_desc}",
                "Org. svinn %": _bl.scrap_pct,
                "Nytt svinn %": _ny_svinn,
                "Org. co-prod %": _bl.co_product_pct,
                "Ny co-prod %": _ny_co,
                "Co-prod vare": _bl.co_product_item_no,
            })
        if _rows:
            _df = pd.DataFrame(_rows)
            bom_scrap_df = mo.ui.data_editor(_df, editable_columns=["Nytt svinn %", "Ny co-prod %"])

    bom_scrap_df
    return (bom_scrap_df,)


@app.cell
def _(bom_scrap_df, persistent_overrides):
    if bom_scrap_df is not None and bom_scrap_df.value is not None:
        _df = bom_scrap_df.value
        for _, _row in _df.iterrows():
            # Strip beskrivelse fra concat-verdier for å gjenopprette original nøkkel
            _komponent = str(_row["Komponent"]).split(" · ")[0]
            _produkt = str(_row["Produkt"]).split(" · ")[0]
            _key = (_produkt, _komponent)

            # Svinn
            _org_svinn = _row["Org. svinn %"]
            _ny_svinn = _row["Nytt svinn %"]
            if abs(_ny_svinn - _org_svinn) > 0.001:
                persistent_overrides["bom_scrap"][_key] = _ny_svinn
            elif _key in persistent_overrides["bom_scrap"]:
                del persistent_overrides["bom_scrap"][_key]

            # Co-prod %
            _org_co = _row["Org. co-prod %"]
            _ny_co = _row["Ny co-prod %"]
            if abs(_ny_co - _org_co) > 0.001:
                persistent_overrides["bom_co_product"][_key] = _ny_co
            elif _key in persistent_overrides["bom_co_product"]:
                del persistent_overrides["bom_co_product"][_key]
    return


@app.cell
def _(mo):
    mo.md("""
    **🏭 Arbeidssentre (timekostnad)**
    """)
    return


@app.cell
def _(data, filtered_work_centers, mo, pd, persistent_overrides):
    # Redigerbar tabell for arbeidssentre
    wc_cost_df = None
    if data:
        _rows = []
        for _wc in filtered_work_centers:
            _saved = persistent_overrides["work_centers"].get(_wc.code, {})
            _ny_lønn = _saved.get("labor_cost_hour", _wc.labor_cost_hour)
            _ny_maskin = _saved.get("machine_cost_hour", _wc.machine_cost_hour)
            _ny_overhead = _saved.get("overhead_cost_hour", _wc.overhead_cost_hour)
            _ny_eff = _saved.get("effective_capacity_pct", _wc.effective_capacity_pct)
            _rows.append({
                "Kode": _wc.code,
                "Beskrivelse": f"{_wc.description} · {_wc.location_code}",
                "Org. lønn": _wc.labor_cost_hour,
                "Ny lønn": _ny_lønn,
                "Org. maskin": _wc.machine_cost_hour,
                "Ny maskin": _ny_maskin,
                "Org. overhead": _wc.overhead_cost_hour,
                "Ny overhead": _ny_overhead,
                "Org. eff. %": _wc.effective_capacity_pct,
                "Ny eff. %": _ny_eff,
            })
        if _rows:
            _df = pd.DataFrame(_rows)
            wc_cost_df = mo.ui.data_editor(_df, editable_columns=["Ny lønn", "Ny maskin", "Ny overhead", "Ny eff. %"])

    wc_cost_df
    return (wc_cost_df,)


@app.cell
def _(persistent_overrides, wc_cost_df):
    if wc_cost_df is not None and wc_cost_df.value is not None:
        _df = wc_cost_df.value
        for _, _row in _df.iterrows():
            _kode = _row["Kode"]
            _wc_overrides = {}
            if abs(_row["Ny lønn"] - _row["Org. lønn"]) > 0.001:
                _wc_overrides["labor_cost_hour"] = _row["Ny lønn"]
            if abs(_row["Ny maskin"] - _row["Org. maskin"]) > 0.001:
                _wc_overrides["machine_cost_hour"] = _row["Ny maskin"]
            if abs(_row["Ny overhead"] - _row["Org. overhead"]) > 0.001:
                _wc_overrides["overhead_cost_hour"] = _row["Ny overhead"]
            if abs(_row["Ny eff. %"] - _row["Org. eff. %"]) > 0.001:
                _wc_overrides["effective_capacity_pct"] = _row["Ny eff. %"]

            if _wc_overrides:
                persistent_overrides["work_centers"][_kode] = _wc_overrides
            elif _kode in persistent_overrides["work_centers"]:
                del persistent_overrides["work_centers"][_kode]
    return


@app.cell
def _(mo):
    mo.md("""
    **📋 Routing (stykkpris)**
    """)
    return


@app.cell
def _(
    data,
    filtered_operations,
    filtered_routing_lines,
    filtered_work_centers,
    mo,
    pd,
    persistent_overrides,
):
    # Redigerbar tabell for routing
    routing_df = None
    if data:
        _rows = []
        for _rl in filtered_routing_lines:
            _op = next((o for o in filtered_operations if o.code == _rl.operation_code), None)
            _op_desc = _op.description if _op else _rl.operation_code
            _wc = next((w for w in filtered_work_centers if w.code == _rl.work_center_code), None)
            _loc = _wc.location_code if _wc else ""

            # Use persistent overrides if exists
            _key = (_rl.item_no, _rl.operation_no, _rl.work_center_code)
            _saved = persistent_overrides["routing"].get(_key, {})
            _ny_run_time = _saved.get("run_time_minutes", _rl.run_time_minutes)
            _ny_batch = _saved.get("batch_size", _rl.batch_size)
            _ny_setup = _saved.get("setup_time_minutes", _rl.setup_time_minutes)

            _rows.append({
                "Produkt": _rl.item_no,
                "Operasjon": f"{_rl.operation_code} ({_op_desc})",
                "Arb.senter": _rl.work_center_code,
                "Lokasjon": _loc,
                "Org. run time": _rl.run_time_minutes,
                "Ny run time": _ny_run_time,
                "Org. setup (min)": _rl.setup_time_minutes,
                "Ny setup (min)": _ny_setup,
                "Org. batch": _rl.batch_size,
                "Ny batch": _ny_batch,
                "_operation_no": _rl.operation_no, # hidden key helper
            })
        if _rows:
            _df = pd.DataFrame(_rows)
            routing_df = mo.ui.data_editor(_df, editable_columns=["Ny run time", "Ny setup (min)", "Ny batch"])

    routing_df
    return (routing_df,)


@app.cell
def _(persistent_overrides, routing_df):
    if routing_df is not None and routing_df.value is not None:
        _df = routing_df.value
        for _, _row in _df.iterrows():
            _produkt = _row["Produkt"]
            _wc = _row["Arb.senter"]
            _op_no = _row["_operation_no"]
            _key = (_produkt, _op_no, _wc)

            _rt_overrides = {}
            if abs(_row["Ny run time"] - _row["Org. run time"]) > 0.001:
                _rt_overrides["run_time_minutes"] = _row["Ny run time"]
            if abs(_row["Ny setup (min)"] - _row["Org. setup (min)"]) > 0.001:
                _rt_overrides["setup_time_minutes"] = _row["Ny setup (min)"]
            if abs(_row["Ny batch"] - _row["Org. batch"]) > 0.001:
                _rt_overrides["batch_size"] = _row["Ny batch"]

            if _rt_overrides:
                persistent_overrides["routing"][_key] = _rt_overrides
            elif _key in persistent_overrides["routing"]:
                del persistent_overrides["routing"][_key]
    return


@app.cell
def _(mo):
    mo.md("""
    **📦 Planlagt kvantum**
    """)
    return


@app.cell
def _(mo):
    mo.md("""
    **Planlagt kvantum** bestemmer hvor mye setupkost som fordeles per enhet.

    Formel: `Setupkost per enhet = Setup-tid (timer) × timekostnad / kvantum`

    Jo høyere kvantum, desto lavere blir setupkostnaden per produsert enhet.
    """)
    return


@app.cell
def _(mo):
    planned_qty = mo.ui.number(
        label="Planlagt kvantum (stk)",
        start=1,
        stop=100000,
        step=1,
        value=1000,
    )
    planned_qty
    return (planned_qty,)


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>🚀 Kjør simulering</h2></div>')
    return


@app.cell
def _(mo):
    run_button = mo.ui.run_button(label="▶️ Kjør simulering", kind="neutral")
    run_button
    return (run_button,)


@app.cell
def _(
    SimulationEngine,
    SimulationOverride,
    data,
    mo,
    pd,
    persistent_overrides,
    planned_qty,
    run_button,
):
    # Initialiser returverdier
    sim_results = None
    sim_overrides = None

    if run_button.value:
        if not data:
            mo.output.replace(mo.md("### ❌ Ingen data lastet. Last opp Excel eller bruk testdata."))
        else:
            try:
                # Bygg overrides fra persistent_overrides
                _overrides = SimulationOverride(
                    planned_quantity=planned_qty.value,
                )

                # Råvarepriser
                for _varenr, _ny_pris in persistent_overrides["item_costs"].items():
                    _overrides.item_costs[_varenr] = _ny_pris

                # Svinn-prosenter
                for _key, _ny_svinn in persistent_overrides["bom_scrap"].items():
                    _overrides.bom_scrap[_key] = _ny_svinn

                # Arbeidssentre
                for _kode, _wc_overrides in persistent_overrides["work_centers"].items():
                    _overrides.work_centers[_kode] = _wc_overrides

                # Routing
                for _key, _rt_overrides in persistent_overrides["routing"].items():
                    _overrides.routing[_key] = _rt_overrides

                # Co-prod %
                for _key, _ny_co in persistent_overrides["bom_co_product"].items():
                    _overrides.bom_co_product[_key] = _ny_co

                _engine = SimulationEngine(data)
                _comparisons = _engine.compare_all(_overrides)

                if _comparisons:
                    sim_results = _comparisons
                    sim_overrides = _overrides

                    _outputs = []

                    # Sammenligningstabell
                    _rows = []
                    for _c in _comparisons:
                        _rows.append({
                            "Lokasjon": _c.location_code,
                            "Produkt": _c.product_no,
                            "Beskrivelse": _c.product_desc,
                            "Org. materialkost": round(_c.original_material_cost, 2),
                            "Sim. materialkost": round(_c.simulated_material_cost, 2),
                            "Diff material": round(_c.material_diff, 2),
                            "Org. operasjonskost": round(_c.original_operation_cost, 2),
                            "Sim. operasjonskost": round(_c.simulated_operation_cost, 2),
                            "Diff operasjon": round(_c.operation_diff, 2),
                            "Org. setupkost": round(_c.original_setup_cost, 2),
                            "Sim. setupkost": round(_c.simulated_setup_cost, 2),
                            "Diff setup": round(_c.setup_diff, 2),
                            "Org. brutto": round(_c.original_gross_cost, 2),
                            "Sim. brutto": round(_c.simulated_gross_cost, 2),
                            "Diff brutto": round(_c.gross_diff, 2),
                            "Org. biprodukt": round(_c.original_byproduct_value, 2),
                            "Sim. biprodukt": round(_c.simulated_byproduct_value, 2),
                            "Diff biprodukt": round(_c.byproduct_diff, 2),
                            "Org. netto": round(_c.original_net_cost, 2),
                            "Sim. netto": round(_c.simulated_net_cost, 2),
                            "Diff netto": round(_c.net_diff, 2),
                        })
                    _df = pd.DataFrame(_rows)
                    _outputs.append(mo.md("**📊 Sammenligning: Baseline vs Simulert**"))
                    _outputs.append(mo.ui.table(_df, selection=None))

                    # Scenario-totaler (per lokasjon)
                    if _comparisons[0].planned_quantity:
                        _sc_rows = []
                        for _c in _comparisons:
                            _sc_rows.append({
                                "Lokasjon": _c.location_code,
                                "Produkt": _c.product_no,
                                "Kvantum": f"{_c.planned_quantity:.0f}",
                                "Total netto kost": f"{_c.simulated_total_net_cost:,.2f}",
                                "Kost per enhet": f"{_c.simulated_cost_per_unit:.2f}",
                                "Timebehov": f"{_c.simulated_total_hours:.2f} timer",
                            })
                        _sc_df = pd.DataFrame(_sc_rows)
                        _outputs.append(mo.md("**📦 Scenariototaler**"))
                        _outputs.append(mo.ui.table(_sc_df, selection=None))

                    # Detaljer per produkt vises ikke i appen - de er tilgjengelig i Excel-eksporten

                    mo.output.replace(mo.vstack(_outputs))
                else:
                    mo.output.replace(mo.md("### Ingen resultater funnet"))

            except Exception as _e:
                mo.output.replace(mo.md(f"### ❌ Feil ved simulering: {_e}"))
                import traceback as _traceback
                _traceback.print_exc()
    return sim_overrides, sim_results


@app.cell
def _(mo):
    mo.Html('<div class="fti-card"><h2>💾 Eksporter resultater</h2><p style="color:#2C5F8A;">Last ned simuleringsresultatene som PDF-rapport eller Excel-fil</p></div>')
    return


@app.cell
def _(mo):
    export_pdf_button = mo.ui.run_button(label="📄 Eksporter til PDF")
    pdf_kommentar = mo.ui.text_area(
        label="Ledelsessammendrag / Kommentar (vises øverst i PDF)",
        placeholder="Skriv f.eks. anbefalinger eller en kort analyse av simuleringen her...",
        rows=4,
    )
    pdf_inkluder_detaljer = mo.ui.checkbox(
        label="Inkluder tekniske detaljer (Material- og operasjonsdetaljer per produkt) i PDF-en",
        value=False,
    )
    mo.vstack([
        mo.md("### 📊 Tilpass PDF-rapport for ledelsen"),
        pdf_kommentar,
        pdf_inkluder_detaljer,
        export_pdf_button,
    ])
    return export_pdf_button, pdf_inkluder_detaljer, pdf_kommentar


@app.cell
def _(
    export_pdf_button,
    generer_rapport,
    mo,
    os,
    pdf_inkluder_detaljer,
    pdf_kommentar,
    registrer_fonter,
    sim_overrides,
    sim_results,
    tempfile,
):
    if export_pdf_button.value:
        try:
            if not sim_results:
                mo.output.replace(mo.md("### ❌ Ingen simuleringsresultater tilgjengelig. Kjør en simulering først."))
            else:
                # Bygg data for PDF-rapporten
                _sammenligninger = []
                for _c in sim_results:
                    _mat_det = []
                    if _c.simulated_material_details:
                        for _md in _c.simulated_material_details:
                            _mat_det.append({
                                "komponent": _md.component,
                                "qty_per": _md.quantity_per,
                                "scrap_pct": _md.scrap_pct,
                                "enhetskost": _md.unit_cost,
                                "total_kost": _md.total_cost,
                            })
                    _op_det = []
                    if _c.simulated_operation_details:
                        for _od in _c.simulated_operation_details:
                            _op_det.append({
                                "operasjon": _od.operation_no,
                                "arbeidssenter": _od.work_center,
                                "run_time": _od.run_time_min,
                                "batch": _od.batch_size,
                                "kost_per_time": _od.cost_per_hour,
                                "total_kost": _od.total_cost,
                            })
                    _bp_det = []
                    if _c.simulated_byproduct_details:
                        for _bd in _c.simulated_byproduct_details:
                            _bp_det.append({
                                "biprodukt": _bd.item_no,
                                "kvantum": _bd.quantity,
                                "markedsverdi": _bd.market_value,
                                "total_verdi": _bd.total_value,
                            })
                    _co_det = []
                    if _c.simulated_co_product_results:
                        for _co in _c.simulated_co_product_results:
                            _co_det.append({
                                "produkt": _co.product_no,
                                "beskrivelse": _co.product_desc,
                                "materialkost": _co.material_cost,
                                "operasjonskost": _co.operation_cost,
                                "setupkost": _co.setup_cost,
                                "brutto": _co.gross_production_cost,
                                "biproduktverdi": _co.by_product_value,
                                "netto": _co.net_production_cost,
                            })
                    _sammenligninger.append({
                        "produkt": _c.product_no,
                        "beskrivelse": _c.product_desc,
                        "org_netto": _c.original_net_cost,
                        "sim_netto": _c.simulated_net_cost,
                        "diff_netto": _c.net_diff,
                        "materialdetaljer": _mat_det,
                        "operasjonsdetaljer": _op_det,
                        "biprodukter": _bp_det,
                        "coprodukter": _co_det,
                        "kvantum": _c.planned_quantity,
                        "total_netto": _c.simulated_total_net_cost,
                        "kost_per_enhet": _c.simulated_cost_per_unit,
                        "timebehov": _c.simulated_total_hours,
                    })

                _overrides_dict = {}
                if sim_overrides:
                    _overrides_dict = {
                        "item_costs": getattr(sim_overrides, 'item_costs', {}),
                        "bom_scrap": dict(getattr(sim_overrides, 'bom_scrap', {})),
                        "work_centers": getattr(sim_overrides, 'work_centers', {}),
                        "routing": dict(getattr(sim_overrides, 'routing', {})),
                        "planned_quantity": getattr(sim_overrides, 'planned_quantity', None),
                    }

                # Beregn kapasitetsbehov per arbeidssenter
                _wc_hours = {}
                for _c in sim_results:
                    if _c.simulated_operation_details and _c.planned_quantity:
                        for _od in _c.simulated_operation_details:
                            _wc = _od.work_center
                            _run_hrs = (_od.run_time_min / 60.0) * _c.planned_quantity
                            _setup_hrs = (_od.setup_time_min / 60.0) * (_c.planned_quantity / _od.batch_size) if _od.batch_size else 0.0
                            _wc_hours[_wc] = _wc_hours.get(_wc, 0.0) + (_run_hrs + _setup_hrs)

                # Generer PDF
                registrer_fonter()
                _output_path = os.path.join(tempfile.gettempdir(), "simuleringsrapport.pdf")
                generer_rapport(
                    _sammenligninger,
                    _overrides_dict,
                    _output_path,
                    tittel="Produksjonskost Simulator - Simuleringsrapport",
                    kommentar=pdf_kommentar.value,
                    inkluder_detaljer=pdf_inkluder_detaljer.value,
                    kapasitet_data=_wc_hours,
                )

                # Tilby nedlasting
                with open(_output_path, "rb") as _f:
                    _pdf_content = _f.read()

                mo.output.replace(
                    mo.vstack([
                        mo.md(f"### ✅ PDF-rapport generert!"),
                        mo.download(
                            label="📥 Last ned PDF-rapport",
                            filename="simuleringsrapport.pdf",
                            data=_pdf_content,
                        ),
                    ])
                )

        except Exception as _e:
            mo.output.replace(mo.md(f"### ❌ Feil ved generering av PDF: {_e}"))
            import traceback as _traceback
            _traceback.print_exc()
    return


@app.cell
def _(mo):
    export_excel_button = mo.ui.run_button(label="💾 Lagre til Excel", kind="neutral")
    export_excel_button
    return (export_excel_button,)


@app.cell
def _(export_excel_button, mo, os, sim_results, tempfile):
    if export_excel_button.value:
        if not sim_results:
            mo.output.replace(mo.md("### ❌ Ingen simuleringsresultater tilgjengelig. Kjør en simulering først."))
        else:
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
                from openpyxl.utils import get_column_letter
                from openpyxl.drawing.image import Image as XLImage
                from PIL import Image as PILImage
                from generer_simuleringsrapport import _hent_logo

                _wb = openpyxl.Workbook()
                _logo = _hent_logo()

                # ── Stiler ──────────────────────────────────────────
                _header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
                _header_fill = PatternFill(start_color='14532D', end_color='14532D', fill_type='solid')
                _subheader_fill = PatternFill(start_color='2F855A', end_color='2F855A', fill_type='solid')
                _title_font = Font(name='Calibri', bold=True, size=14, color='14532D')
                _section_font = Font(name='Calibri', bold=True, size=12, color='2F855A')
                _data_font = Font(name='Calibri', size=10)
                _bold_font = Font(name='Calibri', bold=True, size=10)
                _green_font = Font(name='Calibri', size=10, color='2F855A')
                _red_font = Font(name='Calibri', size=10, color='C53030')
                _thin_border = Border(
                    left=Side(style='thin', color='48BB78'),
                    right=Side(style='thin', color='48BB78'),
                    top=Side(style='thin', color='48BB78'),
                    bottom=Side(style='thin', color='48BB78'),
                )
                _center_align = Alignment(horizontal='center', vertical='center')
                _left_align = Alignment(horizontal='left', vertical='center')
                _right_align = Alignment(horizontal='right', vertical='center')
                _num_fmt = '#,##0.00'
                _pct_fmt = '0.0%'

                def _style_header_row(ws, row, max_col, fill=None):
                    if fill is None:
                        fill = _header_fill
                    for col in range(1, max_col + 1):
                        cell = ws.cell(row=row, column=col)
                        cell.font = _header_font
                        cell.fill = fill
                        cell.alignment = _center_align
                        cell.border = _thin_border

                def _style_data_cell(ws, row, col, is_number=False, is_pct=False):
                    cell = ws.cell(row=row, column=col)
                    cell.font = _data_font
                    cell.border = _thin_border
                    if is_number:
                        cell.alignment = _right_align
                        cell.number_format = _num_fmt
                    elif is_pct:
                        cell.alignment = _right_align
                        cell.number_format = _pct_fmt
                    else:
                        cell.alignment = _left_align
                    return cell

                def _auto_width(ws, max_col, min_width=10, max_width=40):
                    for col in range(1, max_col + 1):
                        letter = get_column_letter(col)
                        max_len = min_width
                        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
                            for cell in row:
                                if cell.value:
                                    max_len = max(max_len, min(len(str(cell.value)), max_width))
                        ws.column_dimensions[letter].width = max_len + 2

                def _sett_logo(ws):
                    """Sett inn Fram Treindustri-logo øverst til venstre i arket."""
                    if _logo and os.path.exists(_logo):
                        try:
                            _img = XLImage(_logo)
                            _img.width = 180
                            _img.height = 45
                            ws.add_image(_img, 'A1')
                            ws.row_dimensions[1].height = 50
                        except Exception:
                            pass

                # ════════════════════════════════════════════════════
                #  ARK 1: SAMMENLIGNING
                # ════════════════════════════════════════════════════
                _ws1 = _wb.active
                _ws1.title = "Sammenligning"

                # Sett inn logo
                _sett_logo(_ws1)

                # Tittel (flyttet til rad 2 pga. logo)
                _ws1.merge_cells('A2:U2')
                _ws1.cell(row=2, column=1, value="Simuleringsresultater - Sammenligning Baseline vs Simulert").font = _title_font
                _ws1.row_dimensions[2].height = 30

                # Overskrifter
                _headers = [
                    "Lokasjon", "Produkt", "Beskrivelse",
                    "Org. materialkost", "Sim. materialkost", "Diff material",
                    "Org. operasjonskost", "Sim. operasjonskost", "Diff operasjon",
                    "Org. setupkost", "Sim. setupkost", "Diff setup",
                    "Org. brutto", "Sim. brutto", "Diff brutto",
                    "Org. biprodukt", "Sim. biprodukt", "Diff biprodukt",
                    "Org. netto", "Sim. netto", "Diff netto",
                ]
                _row_num = 3
                for col_idx, h in enumerate(_headers, 1):
                    _ws1.cell(row=_row_num, column=col_idx, value=h)
                _style_header_row(_ws1, _row_num, len(_headers))

                # Data
                for _c in sim_results:
                    _row_num += 1
                    _data = [
                        _c.location_code, _c.product_no, _c.product_desc,
                        round(_c.original_material_cost, 2), round(_c.simulated_material_cost, 2), round(_c.material_diff, 2),
                        round(_c.original_operation_cost, 2), round(_c.simulated_operation_cost, 2), round(_c.operation_diff, 2),
                        round(_c.original_setup_cost, 2), round(_c.simulated_setup_cost, 2), round(_c.setup_diff, 2),
                        round(_c.original_gross_cost, 2), round(_c.simulated_gross_cost, 2), round(_c.gross_diff, 2),
                        round(_c.original_byproduct_value, 2), round(_c.simulated_byproduct_value, 2), round(_c.byproduct_diff, 2),
                        round(_c.original_net_cost, 2), round(_c.simulated_net_cost, 2), round(_c.net_diff, 2),
                    ]
                    for col_idx, val in enumerate(_data, 1):
                        _is_num = col_idx >= 4  # kolonne 4+ er tall
                        cell = _style_data_cell(_ws1, _row_num, col_idx, is_number=_is_num)
                        cell.value = val
                        # Fargelegg diff-kolonner
                        if col_idx in (6, 9, 12, 15, 18, 21) and isinstance(val, (int, float)):
                            if val > 0:
                                cell.font = _red_font
                            elif val < 0:
                                cell.font = _green_font

                _auto_width(_ws1, len(_headers))

                # ════════════════════════════════════════════════════
                #  ARK 2: SCENARIOTOTALER
                # ════════════════════════════════════════════════════
                _ws2 = _wb.create_sheet("Scenariototaler")
                _sett_logo(_ws2)

                _ws2.merge_cells('A2:G2')
                _ws2.cell(row=2, column=1, value="Scenariototaler").font = _title_font
                _ws2.row_dimensions[2].height = 30

                _sc_headers = ["Lokasjon", "Produkt", "Kvantum", "Total netto kost", "Kost per enhet", "Timebehov"]
                _row_num = 3
                for col_idx, h in enumerate(_sc_headers, 1):
                    _ws2.cell(row=_row_num, column=col_idx, value=h)
                _style_header_row(_ws2, _row_num, len(_sc_headers))

                for _c in sim_results:
                    if _c.planned_quantity:
                        _row_num += 1
                        _data = [
                            _c.location_code, _c.product_no,
                            _c.planned_quantity,
                            _c.simulated_total_net_cost,
                            _c.simulated_cost_per_unit,
                            _c.simulated_total_hours,
                        ]
                        for col_idx, val in enumerate(_data, 1):
                            _is_num = col_idx >= 3
                            cell = _style_data_cell(_ws2, _row_num, col_idx, is_number=_is_num)
                            cell.value = val

                _auto_width(_ws2, len(_sc_headers))

                # ════════════════════════════════════════════════════
                #  ARK 3: CO-PRODUKTER
                # ════════════════════════════════════════════════════
                _ws3 = _wb.create_sheet("Co-produkter")
                _sett_logo(_ws3)

                _ws3.merge_cells('A2:H2')
                _ws3.cell(row=2, column=1, value="Co-produkter (B-vare)").font = _title_font
                _ws3.row_dimensions[2].height = 30

                _co_headers = ["Hovedprodukt", "Co-produkt", "Beskrivelse", "Materialkost", "Operasjonskost", "Setupkost", "Brutto", "Biproduktverdi", "Netto"]
                _row_num = 3
                for col_idx, h in enumerate(_co_headers, 1):
                    _ws3.cell(row=_row_num, column=col_idx, value=h)
                _style_header_row(_ws3, _row_num, len(_co_headers))

                for _c in sim_results:
                    if _c.simulated_co_product_results:
                        for _co in _c.simulated_co_product_results:
                            _row_num += 1
                            _data = [
                                _c.product_no,
                                _co.product_no, _co.product_desc,
                                _co.material_cost, _co.operation_cost, _co.setup_cost,
                                _co.gross_production_cost, _co.by_product_value, _co.net_production_cost,
                            ]
                            for col_idx, val in enumerate(_data, 1):
                                _is_num = col_idx >= 4
                                cell = _style_data_cell(_ws3, _row_num, col_idx, is_number=_is_num)
                                cell.value = val

                _auto_width(_ws3, len(_co_headers))

                # ════════════════════════════════════════════════════
                #  ARK 4: BIPRODUKTER
                # ════════════════════════════════════════════════════
                _ws4 = _wb.create_sheet("Biprodukter")
                _sett_logo(_ws4)

                _ws4.merge_cells('A2:F2')
                _ws4.cell(row=2, column=1, value="Biprodukter").font = _title_font
                _ws4.row_dimensions[2].height = 30

                _bp_headers = ["Produkt", "Biprodukt", "Beskrivelse", "Kvantum", "Markedsverdi", "Total verdi"]
                _row_num = 3
                for col_idx, h in enumerate(_bp_headers, 1):
                    _ws4.cell(row=_row_num, column=col_idx, value=h)
                _style_header_row(_ws4, _row_num, len(_bp_headers))

                for _c in sim_results:
                    if _c.simulated_byproduct_details:
                        for _bd in _c.simulated_byproduct_details:
                            _row_num += 1
                            _data = [
                                _c.product_no,
                                _bd.item_no, _bd.description,
                                _bd.quantity, _bd.market_value, _bd.total_value,
                            ]
                            for col_idx, val in enumerate(_data, 1):
                                _is_num = col_idx >= 4
                                cell = _style_data_cell(_ws4, _row_num, col_idx, is_number=_is_num)
                                cell.value = val

                _auto_width(_ws4, len(_bp_headers))

                # ════════════════════════════════════════════════════
                #  ARK 5: DETALJER PER PRODUKT
                # ════════════════════════════════════════════════════
                _ws5 = _wb.create_sheet("Detaljer")
                _sett_logo(_ws5)

                _ws5.merge_cells('A2:J2')
                _ws5.cell(row=2, column=1, value="Detaljer per produkt").font = _title_font
                _ws5.row_dimensions[2].height = 30

                _row_num = 3
                for _c in sim_results:
                    # Produktoverskrift
                    _ws5.merge_cells(f'A{_row_num}:J{_row_num}')
                    _ws5.cell(row=_row_num, column=1, value=f"{_c.product_no} - {_c.product_desc} ({_c.location_code})").font = _section_font
                    _ws5.row_dimensions[_row_num].height = 22
                    _row_num += 1

                    # Materialdetaljer
                    if _c.simulated_material_details:
                        _mat_headers = ["Komponent", "Beskrivelse", "Qty per", "Svinn %", "Enhetskost", "Total kost"]
                        for col_idx, h in enumerate(_mat_headers, 1):
                            _ws5.cell(row=_row_num, column=col_idx, value=h)
                        _style_header_row(_ws5, _row_num, len(_mat_headers), fill=_subheader_fill)
                        _row_num += 1

                        for _md in _c.simulated_material_details:
                            _data = [
                                _md.component, _md.component_desc,
                                _md.quantity_per, _md.scrap_pct / 100.0,
                                _md.unit_cost, _md.total_cost,
                            ]
                            for col_idx, val in enumerate(_data, 1):
                                _is_num = col_idx >= 3
                                _is_pct = col_idx == 4
                                cell = _style_data_cell(_ws5, _row_num, col_idx, is_number=_is_num, is_pct=_is_pct)
                                cell.value = val
                            _row_num += 1

                    # Operasjonsdetaljer
                    if _c.simulated_operation_details:
                        _op_headers = ["Op.nr", "Beskrivelse", "Arbeidssenter", "Run time (min)", "Setup (min)", "Batch", "Kost/time", "Run kost", "Setup/unit", "Total"]
                        for col_idx, h in enumerate(_op_headers, 1):
                            _ws5.cell(row=_row_num, column=col_idx, value=h)
                        _style_header_row(_ws5, _row_num, len(_op_headers), fill=_subheader_fill)
                        _row_num += 1

                        for _od in _c.simulated_operation_details:
                            _data = [
                                _od.operation_no, _od.operation_desc, _od.work_center,
                                _od.run_time_min, _od.setup_time_min, _od.batch_size,
                                _od.cost_per_hour, _od.run_cost, _od.setup_cost_per_unit, _od.total_cost,
                            ]
                            for col_idx, val in enumerate(_data, 1):
                                _is_num = col_idx >= 4
                                cell = _style_data_cell(_ws5, _row_num, col_idx, is_number=_is_num)
                                cell.value = val
                            _row_num += 1

                    _row_num += 1  # Tom rad mellom produkter

                _auto_width(_ws5, 10)

                # ── Lagre til temp-fil ──────────────────────────────
                _temp_path = os.path.join(tempfile.gettempdir(), "simuleringsresultater.xlsx")
                _wb.save(_temp_path)

                with open(_temp_path, "rb") as _f:
                    _excel_data = _f.read()

                mo.output.replace(
                    mo.vstack([
                        mo.md("### ✅ Excel-rapport generert!"),
                        mo.download(
                            label="📥 Last ned Excel-rapport",
                            filename="simuleringsresultater.xlsx",
                            data=_excel_data,
                        )
                    ])
                )
            except ImportError:
                mo.output.replace(mo.md("### ❌ openpyxl er ikke installert. Kjør: pip install openpyxl"))
            except Exception as _e:
                mo.output.replace(mo.md(f"### ❌ Feil ved generering av Excel: {_e}"))
                import traceback as _traceback
                _traceback.print_exc()
    return


if __name__ == "__main__":
    app.run()
