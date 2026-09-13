#!/usr/bin/env python3
"""Smoke test for PostgreSQL-backend.

Testen er laget for lokal bruk og CI. Den verifiserer at PostgreSQL-sporet kan:

1. initialisere skjemaet
2. skrive og lese via DataRepo
3. lese de samme dataene via DatabaseData (kalkylemotorens datakilde)
4. importere Excel til PostgreSQL og eksportere PostgreSQL tilbake til Excel
5. rydde opp testdata etter seg

Bruk (PowerShell):
    $env:DATABASE_URL="postgresql://user:pass@host:5432/db?sslmode=require"
    python src/scripts/postgres_smoke_test.py
"""

import os
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

_src = Path(__file__).resolve().parents[1]
if str(_src) not in sys.path:
    sys.path.insert(0, str(_src))

from data_repo import DataRepo
from excel_bridge import export_sqlite_to_excel, import_excel_to_sqlite
from kostberegning import DatabaseData


SMOKE_ITEM = "SMOKE_FG_001"
EXCEL_SMOKE_ITEM = "SMOKE_XLSX_001"
EXCEL_SMOKE_FILENAME = "postgres_smoke_import.xlsx"


def _cleanup(db: DataRepo, row_ids: list[int] | None = None) -> None:
    """Fjern testdata uten å bruke CRUD-logging."""
    db.execute(
        "DELETE FROM transport_flagg WHERE item_no IN (?, ?)",
        (SMOKE_ITEM, EXCEL_SMOKE_ITEM),
    )
    db.execute(
        "DELETE FROM products WHERE item_no IN (?, ?)",
        (SMOKE_ITEM, EXCEL_SMOKE_ITEM),
    )
    db.execute("DELETE FROM uploaded_files WHERE filename = ?", (EXCEL_SMOKE_FILENAME,))
    for row_id in row_ids or []:
        db.execute("DELETE FROM change_log WHERE record_key = ?", (str(row_id),))
    db.conn.commit()


def _make_excel_fixture(path: Path) -> None:
    """Lag en minimal, gyldig Excel-importfil for Product Master."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Product Master"
    ws.append([
        "ACTION",
        "Rad ID",
        "Item No",
        "Description",
        "Item Type",
        "Product Group",
        "Base Unit of Measure",
        "Is Transport",
    ])
    ws.append([
        "CREATE",
        None,
        EXCEL_SMOKE_ITEM,
        "Excel smoke test product",
        "Finished Good",
        "TEST",
        "LM",
        1,
    ])
    wb.save(path)


def _assert_excel_export(path: Path) -> None:
    """Verifiser at eksportert arbeidsbok inneholder Excel-testproduktet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        product_sheet = None
        for ws in wb.worksheets:
            # Dataark med ACTION/Rad ID har Item No i kolonne C.
            if ws.cell(row=1, column=3).value == "Item No":
                product_sheet = ws
                break
        if product_sheet is None:
            raise AssertionError("Fant ikke Product Master-ark i eksportert Excel-fil")

        exported_items = {
            str(product_sheet.cell(row=row, column=3).value or "").strip()
            for row in range(2, product_sheet.max_row + 1)
        }
        if EXCEL_SMOKE_ITEM not in exported_items:
            raise AssertionError(
                f"Eksportert Excel mangler testproduktet {EXCEL_SMOKE_ITEM}"
            )
    finally:
        wb.close()


def main() -> int:
    database_url = os.environ.get("DATABASE_URL", "").strip()
    if not database_url:
        print("[ERROR] DATABASE_URL er ikke satt.")
        return 1

    db = DataRepo()
    row_ids: list[int] = []
    data: DatabaseData | None = None

    try:
        if db.backend != "postgresql":
            print(f"[ERROR] Forventet PostgreSQL-backend, fikk: {db.backend}")
            return 2

        print(f"[INFO] Backend: {db.backend}")
        db.initialize()
        print("[OK] initialize() fullført")

        # Gjør testen idempotent etter eventuelle avbrutte tidligere kjøringer.
        _cleanup(db)

        ids = db.upsert_products([
            {
                "item_no": SMOKE_ITEM,
                "description": "Smoke test product",
                "item_type": "FG",
                "product_group": "TEST",
                "base_uom": "LM",
            }
        ], source="smoke_test")
        smoke_id = int(ids[0]) if ids else None
        if smoke_id is not None:
            row_ids.append(smoke_id)
        print("[OK] upsert_products() fullført")

        smoke_row = db.execute(
            "SELECT id, item_no, description FROM products WHERE item_no = ?",
            (SMOKE_ITEM,),
        ).fetchone()
        if smoke_row is None:
            print("[ERROR] Fant ikke smoke test-produkt etter insert.")
            return 3

        smoke_id = int(smoke_row["id"])
        if smoke_id not in row_ids:
            row_ids.append(smoke_id)
        print(f"[OK] DataRepo roundtrip: {smoke_row['item_no']} / {smoke_row['description']}")

        # Verifiser den faktiske datakilden som CostCalculator bruker i appen.
        data = DatabaseData()
        product = data.product(SMOKE_ITEM)
        if product is None:
            print("[ERROR] DatabaseData fant ikke smoke test-produktet.")
            return 4
        if product.description != "Smoke test product":
            print(f"[ERROR] Uventet produktbeskrivelse: {product.description}")
            return 5
        print("[OK] DatabaseData leser fra PostgreSQL")

        # Verifiser hele Excel-broen mot PostgreSQL, ikke bare DataRepo.
        with TemporaryDirectory(prefix="prodcalc-postgres-smoke-") as tmpdir:
            tmp = Path(tmpdir)
            import_path = tmp / EXCEL_SMOKE_FILENAME
            export_path = tmp / "postgres_smoke_export.xlsx"
            _make_excel_fixture(import_path)

            import_stats = import_excel_to_sqlite(
                str(import_path),
                db=db,
                excel_blob=import_path.read_bytes(),
                comment="PostgreSQL CI smoke test",
            )
            if import_stats.get("errors"):
                print(f"[ERROR] Excel-import feilet: {import_stats['errors']}")
                return 6

            excel_row = db.execute(
                "SELECT id, description FROM products WHERE item_no = ?",
                (EXCEL_SMOKE_ITEM,),
            ).fetchone()
            if excel_row is None:
                print("[ERROR] Excel-import opprettet ikke testproduktet i PostgreSQL.")
                return 7
            excel_id = int(excel_row["id"])
            row_ids.append(excel_id)
            if excel_row["description"] != "Excel smoke test product":
                print(f"[ERROR] Uventet Excel-importert beskrivelse: {excel_row['description']}")
                return 8

            transport_row = db.execute(
                "SELECT is_transport FROM transport_flagg WHERE item_no = ?",
                (EXCEL_SMOKE_ITEM,),
            ).fetchone()
            if transport_row is None or not bool(transport_row["is_transport"]):
                print("[ERROR] Excel-import lagret ikke Is Transport i PostgreSQL.")
                return 9
            print("[OK] Excel → PostgreSQL fungerer")

            export_sqlite_to_excel(str(export_path), db=db)
            _assert_excel_export(export_path)
            print("[OK] PostgreSQL → Excel fungerer")

        print("[OK] PostgreSQL smoke test fullført")
        return 0
    finally:
        # Dersom testen feilet midt i en PostgreSQL-transaksjon må vi rulle
        # tilbake før cleanup, ellers maskerer InFailedSqlTransaction rotfeilen.
        try:
            if db._conn is not None:
                db.conn.rollback()
            _cleanup(db, row_ids)
        finally:
            if data is not None:
                data.db.close()
            db.close()


if __name__ == "__main__":
    raise SystemExit(main())
