#!/usr/bin/env python3
"""
generer_excel_rapport.py - Genererer en Excel-rapport fra simuleringsresultater.

Brukes av varekost_app.py for å eksportere resultater til Excel.
Kan også kjøres frittstående med JSON-data.

Avhengigheter: openpyxl, Pillow (begge er i requirements.txt)
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from generer_simuleringsrapport import _hent_logo


def generer_excel_rapport(sim_results, output_path):
    """
    Generer en Excel-rapport fra simuleringsresultater.

    Args:
        sim_results: Liste med SimulationComparison-objekter
        output_path: Sti til output Excel-fil (.xlsx)

    Returns:
        bytes: Innholdet av Excel-filen (for nedlasting)
    """
    _wb = openpyxl.Workbook()
    _logo = _hent_logo()

    # ── Stiler ────────────────────────────────────────────────────
    _header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    _header_fill = PatternFill(start_color='14532D', end_color='14532D', fill_type='solid')
    _subheader_fill = PatternFill(start_color='2F855A', end_color='2F855A', fill_type='solid')
    _title_font = Font(name='Calibri', bold=True, size=14, color='14532D')
    _section_font = Font(name='Calibri', bold=True, size=12, color='2F855A')
    _data_font = Font(name='Calibri', size=10)
    _green_font = Font(name='Calibri', size=10, color='2F855A')
    _red_font = Font(name='Calibri', size=10, color='C53030')
    _thin_border = Border(
        left=Side(style='thin', color='48BB78'),
        right=Side(style='thin', color='48BB78'),
        top=Side(style='thin', color='48BB78'),
        bottom=Side(style='thin', color='48BB78'),
    )
    _center_align = Alignment(horizontal='center', vertical='center')
    _left_align = Alignment(horizontal='left', vertical='center')
    _right_align = Alignment(horizontal='right', vertical='center')
    _num_fmt = '#,##0.00'
    _pct_fmt = '0.0%'

    def _style_header_row(ws, row, max_col, fill=None):
        if fill is None:
            fill = _header_fill
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = _header_font
            cell.fill = fill
            cell.alignment = _center_align
            cell.border = _thin_border

    def _style_data_cell(ws, row, col, is_number=False, is_pct=False):
        cell = ws.cell(row=row, column=col)
        cell.font = _data_font
        cell.border = _thin_border
        if is_number:
            cell.alignment = _right_align
            cell.number_format = _num_fmt
        elif is_pct:
            cell.alignment = _right_align
            cell.number_format = _pct_fmt
        else:
            cell.alignment = _left_align
        return cell

    def _auto_width(ws, max_col, min_width=10, max_width=40):
        for col in range(1, max_col + 1):
            letter = get_column_letter(col)
            max_len = min_width
            for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
                for cell in row:
                    if cell.value:
                        max_len = max(max_len, min(len(str(cell.value)), max_width))
            ws.column_dimensions[letter].width = max_len + 2

    def _sett_logo(ws):
        """Sett inn Fram Treindustri-logo øverst til venstre i arket."""
        if _logo and os.path.exists(_logo):
            try:
                _img = XLImage(_logo)
                _img.width = 180
                _img.height = 45
                ws.add_image(_img, 'A1')
                ws.row_dimensions[1].height = 50
            except Exception:
                pass

    # ── ARK 1: SAMMENLIGNING ──────────────────────────────────────
    _ws1 = _wb.active
    _ws1.title = "Sammenligning"
    _sett_logo(_ws1)

    _ws1.merge_cells('A2:U2')
    _ws1.cell(row=2, column=1, value="Simuleringsresultater - Sammenligning Baseline vs Simulert").font = _title_font
    _ws1.row_dimensions[2].height = 30

    _headers = [
        "Lokasjon", "Produkt", "Beskrivelse",
        "Org. materialkost", "Sim. materialkost", "Diff material",
        "Org. operasjonskost", "Sim. operasjonskost", "Diff operasjon",
        "Org. setupkost", "Sim. setupkost", "Diff setup",
        "Org. brutto", "Sim. brutto", "Diff brutto",
        "Org. biprodukt", "Sim. biprodukt", "Diff biprodukt",
        "Org. netto", "Sim. netto", "Diff netto",
    ]
    _row_num = 3
    for col_idx, h in enumerate(_headers, 1):
        _ws1.cell(row=_row_num, column=col_idx, value=h)
    _style_header_row(_ws1, _row_num, len(_headers))

    for _c in sim_results:
        _row_num += 1
        _data = [
            _c.location_code, _c.product_no, _c.product_desc,
            round(_c.original_material_cost, 2), round(_c.simulated_material_cost, 2), round(_c.material_diff, 2),
            round(_c.original_operation_cost, 2), round(_c.simulated_operation_cost, 2), round(_c.operation_diff, 2),
            round(_c.original_setup_cost, 2), round(_c.simulated_setup_cost, 2), round(_c.setup_diff, 2),
            round(_c.original_gross_cost, 2), round(_c.simulated_gross_cost, 2), round(_c.gross_diff, 2),
            round(_c.original_byproduct_value, 2), round(_c.simulated_byproduct_value, 2), round(_c.byproduct_diff, 2),
            round(_c.original_net_cost, 2), round(_c.simulated_net_cost, 2), round(_c.net_diff, 2),
        ]
        for col_idx, val in enumerate(_data, 1):
            _is_num = col_idx >= 4
            cell = _style_data_cell(_ws1, _row_num, col_idx, is_number=_is_num)
            cell.value = val
            if col_idx in (6, 9, 12, 15, 18, 21) and isinstance(val, (int, float)):
                if val > 0:
                    cell.font = _red_font
                elif val < 0:
                    cell.font = _green_font

    _auto_width(_ws1, len(_headers))

    # ── ARK 2: SCENARIOTOTALER ────────────────────────────────────
    _ws2 = _wb.create_sheet("Scenariototaler")
    _sett_logo(_ws2)

    _ws2.merge_cells('A2:G2')
    _ws2.cell(row=2, column=1, value="Scenariototaler").font = _title_font
    _ws2.row_dimensions[2].height = 30

    _sc_headers = ["Lokasjon", "Produkt", "Kvantum", "Total netto kost", "Kost per enhet", "Timebehov"]
    _row_num = 3
    for col_idx, h in enumerate(_sc_headers, 1):
        _ws2.cell(row=_row_num, column=col_idx, value=h)
    _style_header_row(_ws2, _row_num, len(_sc_headers))

    for _c in sim_results:
        if _c.planned_quantity:
            _row_num += 1
            _data = [
                _c.location_code, _c.product_no,
                _c.planned_quantity,
                _c.simulated_total_net_cost,
                _c.simulated_cost_per_unit,
                _c.simulated_total_hours,
            ]
            for col_idx, val in enumerate(_data, 1):
                _is_num = col_idx >= 3
                cell = _style_data_cell(_ws2, _row_num, col_idx, is_number=_is_num)
                cell.value = val

    _auto_width(_ws2, len(_sc_headers))

    # ── ARK 3: CO-PRODUKTER ──────────────────────────────────────
    _ws3 = _wb.create_sheet("Co-produkter")
    _sett_logo(_ws3)

    _ws3.merge_cells('A2:H2')
    _ws3.cell(row=2, column=1, value="Co-produkter (B-vare)").font = _title_font
    _ws3.row_dimensions[2].height = 30

    _co_headers = ["Hovedprodukt", "Co-produkt", "Beskrivelse", "Materialkost", "Operasjonskost", "Setupkost", "Brutto", "Biproduktverdi", "Netto"]
    _row_num = 3
    for col_idx, h in enumerate(_co_headers, 1):
        _ws3.cell(row=_row_num, column=col_idx, value=h)
    _style_header_row(_ws3, _row_num, len(_co_headers))

    for _c in sim_results:
        if _c.simulated_co_product_results:
            for _co in _c.simulated_co_product_results:
                _row_num += 1
                _data = [
                    _c.product_no,
                    _co.product_no, _co.product_desc,
                    _co.material_cost, _co.operation_cost, _co.setup_cost,
                    _co.gross_production_cost, _co.by_product_value, _co.net_production_cost,
                ]
                for col_idx, val in enumerate(_data, 1):
                    _is_num = col_idx >= 4
                    cell = _style_data_cell(_ws3, _row_num, col_idx, is_number=_is_num)
                    cell.value = val

    _auto_width(_ws3, len(_co_headers))

    # ── ARK 4: BIPRODUKTER ───────────────────────────────────────
    _ws4 = _wb.create_sheet("Biprodukter")
    _sett_logo(_ws4)

    _ws4.merge_cells('A2:F2')
    _ws4.cell(row=2, column=1, value="Biprodukter").font = _title_font
    _ws4.row_dimensions[2].height = 30

    _bp_headers = ["Produkt", "Biprodukt", "Beskrivelse", "Kvantum", "Markedsverdi", "Total verdi"]
    _row_num = 3
    for col_idx, h in enumerate(_bp_headers, 1):
        _ws4.cell(row=_row_num, column=col_idx, value=h)
    _style_header_row(_ws4, _row_num, len(_bp_headers))

    for _c in sim_results:
        if _c.simulated_byproduct_details:
            for _bd in _c.simulated_byproduct_details:
                _row_num += 1
                _data = [
                    _c.product_no,
                    _bd.item_no, _bd.description,
                    _bd.quantity, _bd.market_value, _bd.total_value,
                ]
                for col_idx, val in enumerate(_data, 1):
                    _is_num = col_idx >= 4
                    cell = _style_data_cell(_ws4, _row_num, col_idx, is_number=_is_num)
                    cell.value = val

    _auto_width(_ws4, len(_bp_headers))

    # ── ARK 5: DETALJER PER PRODUKT ─────────────────────────────
    _ws5 = _wb.create_sheet("Detaljer")
    _sett_logo(_ws5)

    _ws5.merge_cells('A2:J2')
    _ws5.cell(row=2, column=1, value="Detaljer per produkt").font = _title_font
    _ws5.row_dimensions[2].height = 30

    _row_num = 3
    for _c in sim_results:
        # Produktoverskrift
        _ws5.merge_cells(f'A{_row_num}:J{_row_num}')
        _ws5.cell(row=_row_num, column=1, value=f"{_c.product_no} - {_c.product_desc} ({_c.location_code})").font = _section_font
        _ws5.row_dimensions[_row_num].height = 22
        _row_num += 1

        # Materialdetaljer
        if _c.simulated_material_details:
            _mat_headers = ["Komponent", "Beskrivelse", "Qty per", "Svinn %", "Enhetskost", "Total kost"]
            for col_idx, h in enumerate(_mat_headers, 1):
                _ws5.cell(row=_row_num, column=col_idx, value=h)
            _style_header_row(_ws5, _row_num, len(_mat_headers), fill=_subheader_fill)
            _row_num += 1

            for _md in _c.simulated_material_details:
                _data = [
                    _md.component, _md.component_desc,
                    _md.quantity_per, _md.scrap_pct / 100.0,
                    _md.unit_cost, _md.total_cost,
                ]
                for col_idx, val in enumerate(_data, 1):
                    _is_num = col_idx >= 3
                    _is_pct = col_idx == 4
                    cell = _style_data_cell(_ws5, _row_num, col_idx, is_number=_is_num, is_pct=_is_pct)
                    cell.value = val
                _row_num += 1

        # Operasjonsdetaljer
        if _c.simulated_operation_details:
            _op_headers = ["Op.nr", "Beskrivelse", "Arbeidssenter", "Run time (min)", "Setup (min)", "Batch", "Kost/time", "Run kost", "Setup/unit", "Total"]
            for col_idx, h in enumerate(_op_headers, 1):
                _ws5.cell(row=_row_num, column=col_idx, value=h)
            _style_header_row(_ws5, _row_num, len(_op_headers), fill=_subheader_fill)
            _row_num += 1

            for _od in _c.simulated_operation_details:
                _data = [
                    _od.operation_no, _od.operation_desc, _od.work_center,
                    _od.run_time_min, _od.setup_time_min, _od.batch_size,
                    _od.cost_per_hour, _od.run_cost, _od.setup_cost_per_unit, _od.total_cost,
                ]
                for col_idx, val in enumerate(_data, 1):
                    _is_num = col_idx >= 4
                    cell = _style_data_cell(_ws5, _row_num, col_idx, is_number=_is_num)
                    cell.value = val
                _row_num += 1

        _row_num += 1  # Tom rad mellom produkter

    _auto_width(_ws5, 10)

    # ── Lagre og returner bytes ───────────────────────────────────
    _wb.save(output_path)
    with open(output_path, "rb") as _f:
        return _f.read()


def main():
    """Kjøres frittstående: python generer_excel_rapport.py"""
    import json
    import sys
    from kostberegning import SimulationComparison

    if len(sys.argv) < 2:
        print("Bruk: python generer_excel_rapport.py data.json [rapport.xlsx]")
        return

    with open(sys.argv[1], 'r', encoding='utf-8') as f:
        data = json.load(f)

    output = sys.argv[2] if len(sys.argv) > 2 else "simuleringsresultater.xlsx"

    # Rekonstruer SimulationComparison-objekter fra JSON
    if isinstance(data, list) and len(data) > 0:
        sim_list = []
        for item in data:
            sc = SimulationComparison(**item)
            sim_list.append(sc)
    else:
        sim_list = data

    generer_excel_rapport(sim_list, output)
    print(f"Excel-rapport generert: {output}")


if __name__ == "__main__":
    main()