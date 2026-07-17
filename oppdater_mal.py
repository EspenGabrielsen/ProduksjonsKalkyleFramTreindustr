"""
oppdater_mal.py - Oppretter oppdatert Produksjonsmodell_Mal.xlsx
med kolonnebeskrivelser og statiske dropdowns.

Kryss-referanser valideres av kostberegning.py under innlesing
(ikke i Excel, for a unnga eksterne koblings-feil).
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)

MAX_ROW = 2000

# ── Kolonnebeskrivelser ──────────────────────────────────────────
KOLONNER = {
    "Product Master": {
        "Item No": "Unik identifikator for varen. Eksempel: RM001, FG001, BP001",
        "Description": "Beskrivende navn pa varen. Eksempel: Skrulast 48x198",
        "Item Type": "Type vare: Raw Material, Semi Finished, Finished Good, By Product, Trading Item",
        "Product Group": "Gruppering av varer. Eksempel: Skrulast, Panel, Kledning, Spon",
        "Base Unit of Measure": "Standard maleenhet. Eksempel: LM, M3, KG, PCS",
        "Active": "Er varen aktiv? Ja / Nei",
    },
    "Locations": {
        "Location Code": "Unik kode for lokasjonen. Eksempel: KOD",
        "Location Name": "Navn pa lokasjonen. Eksempel: Kodal Fabrikk",
        "Location Type": "Type lokasjon: Factory, Warehouse, Distribution Center, Sales Office",
        "Active": "Er lokasjonen aktiv? Ja / Nei",
    },
    "Work Centers": {
        "Work Center Code": "Unik identifikator for arbeidssenteret. Eksempel: HOVEDHOVEL",
        "Description": "Beskrivende navn pa arbeidssenteret",
        "Location Code": "Fabrikken arbeidssenteret tilhorer",
        "Labor Cost per Hour": "Arbeidskostnad per time (lonn, avgift, pensjon, ferie). Eksempel: 550",
        "Machine Cost per Hour": "Maskinkostnad per time (avskrivn., service, leasing, vedlikehold, energi). Eksempel: 900",
        "Overhead Cost per Hour": "Indirekte produksjonskostnader. Eksempel: 150",
        "Capacity Hours per Day": "Tilgjengelige timer per dag. Eksempel: 16",
        "Effective Capacity %": "Hvor stor del av tiden som kan brukes til produksjon. Eksempel: 85",
        "Active": "Er arbeidssenteret aktivt? Ja / Nei",
    },
    "Operation Master": {
        "Operation Code": "Unik kode for operasjonen. Eksempel: HOVLING, MALING, PACKING",
        "Description": "Beskrivelse av operasjonen",
        "Default Work Center": "Anbefalt arbeidssenter for operasjonen",
        "Standard Unit": "Maleenhet for produksjonstid. Eksempel: Minutes, Hours",
        "Active": "Er operasjonen aktiv? Ja / Nei",
    },
    "Item Costs": {
        "Item No": "Referanse til varen (Item No fra Product Master)",
        "Cost Type": "Type kostpris: Standard Cost, Last Direct Cost, Forecast Cost, Budget Cost",
        "Unit Cost": "Kostpris per enhet. Eksempel: 3000.00",
        "Currency": "Valuta. Eksempel: NOK, EUR",
        "Effective Date": "Dato kostprisen gjelder fra",
    },
    "BOM": {
        "Parent Item No": "Produktet som produseres (Item No)",
        "Component Item No": "Komponenten som forbrukes (Item No)",
        "Quantity Per": "Antall output-enheter per input-enhet. Eksempel: 400",
        "Unit of Measure": "Maleenhet. Eksempel: LM",
        "Scrap %": "Forventet materialsvinn i prosent. Eksempel: 5",
        "Co-Prod %": "Andel samprodukt (co-product). Eksempel: 6",
        "Co-Prod Item No": "Varenummer for samproduktet. Eksempel: JD16073-B",
        "Valid From": "Gyldig fra dato",
        "Valid To": "Gyldig til dato (tom = alltid)",
    },
    "Routing": {
        "Item No": "Produkt som produseres (Item No)",
        "Operation No": "Sekvensnummer. Eksempel: 10, 20, 30",
        "Operation Code": "Operasjon (ref. Operation Master)",
        "Work Center Code": "Arbeidssenter (ref. Work Centers)",
        "Setup Time Minutes": "Klargjoringstid i minutter. Eksempel: 15",
        "Run Time Minutes": "Produksjonstid per enhet i minutter. Eksempel: 0.15",
        "Batch Size": "Normal ordrestorrelse. Eksempel: 500",
        "Valid From": "Gyldig fra dato",
        "Valid To": "Gyldig til dato (tom = alltid)",
    },
    "By Product Rules": {
        "Parent Item No": "Produktet (ferdigvaren) som skaper biproduktet",
        "By Product Item No": "Biproduktet (Item No). Eksempel: BP001",
        "Expected Quantity": "Forventet mengde biprodukt per enhet ferdigvare",
        "Unit of Measure": "Maleenhet. Eksempel: KG",
        "Market Value": "Forventet markedspris per enhet. Eksempel: 1.50",
        "Allocation Method": "Reduce Main Product Cost, Separate Profit Center, Informational Only",
    },
    "Capacity Calendar": {
        "Work Center": "Arbeidssenter (ref. Work Centers)",
        "Date": "Dato",
        "Available Hours": "Tilgjengelige timer for dagen",
        "Planned Downtime": "Planlagte stopp i timer (vedlikehold, ferie, ombygging)",
        "Available Production Hours": "Timer til produksjon = Available - Planned",
    },
    "Production Scenario": {
        "Scenario Name": "Navn pa scenario. Eksempel: Normal Produksjon",
        "Product": "Produktet som simuleres (Item No)",
        "Planned Quantity": "Planlagt produksjonsmengde. Eksempel: 100000",
        "Start Date": "Startdato for scenario",
        "End Date": "Sluttdato for scenario",
    },
}

# ── Statiske dropdowns ───────────────────────────────────────────
STATISKE_DROPDOWNS = [
    ("Product Master", "C", '"Raw Material,Semi Finished,Finished Good,By Product,Trading Item"'),
    ("Product Master", "F", '"Ja,Nei"'),
    ("Locations", "C", '"Factory,Warehouse,Distribution Center,Sales Office"'),
    ("Locations", "D", '"Ja,Nei"'),
    ("Work Centers", "I", '"Ja,Nei"'),
    ("Operation Master", "D", '"Minutes,Hours"'),
    ("Operation Master", "E", '"Ja,Nei"'),
    ("Item Costs", "B", '"Standard Cost,Last Direct Cost,Forecast Cost,Budget Cost"'),
    ("Item Costs", "D", '"NOK,EUR,USD,SEK,DKK"'),
    ("By Product Rules", "F", '"Reduce Main Product Cost,Separate Profit Center,Informational Only"'),
]


def legg_til_dropdowns(ws, ark_navn):
    """Legg til statiske dropdowns for et ark."""
    for sheet, col, liste in STATISKE_DROPDOWNS:
        if sheet != ark_navn:
            continue
        dv = DataValidation(
            type="list",
            formula1=liste,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Ugyldig verdi",
            error=f"Verdien må være en av: {liste.replace(chr(34), '')}",
        )
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{MAX_ROW}")


def lag_ark(wb, navn, headers):
    """Opprett et ark med overskrifter, kolonnebeskrivelser og dropdowns."""
    ws = wb.create_sheet(navn)
    ws.append(headers)
    desc = KOLONNER.get(navn, {})

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = THIN_BORDER
        if header in desc:
            cell.comment = Comment(desc[header], "System")

    for col in range(1, len(headers) + 1):
        max_len = len(str(headers[col - 1]))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max_len + 4, 35)

    legg_til_dropdowns(ws, navn)
    return ws


# ── Opprett workbook ─────────────────────────────────────────────
wb = openpyxl.Workbook()

# 1. Product Master (manuell fordi det er default arket)
ws = wb.active
ws.title = "Product Master"
headers = ["Item No", "Description", "Item Type", "Product Group", "Base Unit of Measure", "Active"]
ws.append(headers)
desc = KOLONNER["Product Master"]
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='center', wrap_text=True)
    cell.border = THIN_BORDER
    if header in desc:
        cell.comment = Comment(desc[header], "System")
for col in range(1, len(headers) + 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(len(headers[col - 1]) + 4, 35)
legg_til_dropdowns(ws, "Product Master")

# 2–10. Øvrige ark
lag_ark(wb, "Locations", ["Location Code", "Location Name", "Location Type", "Active"])
lag_ark(wb, "Work Centers", [
    "Work Center Code", "Description", "Location Code",
    "Labor Cost per Hour", "Machine Cost per Hour", "Overhead Cost per Hour",
    "Capacity Hours per Day", "Effective Capacity %", "Active",
])
lag_ark(wb, "Operation Master", ["Operation Code", "Description", "Default Work Center", "Standard Unit", "Active"])
lag_ark(wb, "Item Costs", ["Item No", "Cost Type", "Unit Cost", "Currency", "Effective Date"])
lag_ark(wb, "BOM", [
    "Parent Item No", "Component Item No", "Quantity Per",
    "Unit of Measure", "Scrap %", "Co-Prod %", "Co-Prod Item No",
    "Valid From", "Valid To",
])
lag_ark(wb, "Routing", [
    "Item No", "Operation No", "Operation Code", "Work Center Code",
    "Setup Time Minutes", "Run Time Minutes", "Batch Size",
    "Valid From", "Valid To",
])
lag_ark(wb, "By Product Rules", [
    "Parent Item No", "By Product Item No", "Expected Quantity",
    "Unit of Measure", "Market Value", "Allocation Method",
])
lag_ark(wb, "Capacity Calendar", [
    "Work Center", "Date", "Available Hours", "Planned Downtime", "Available Production Hours",
])
lag_ark(wb, "Production Scenario", [
    "Scenario Name", "Product", "Planned Quantity", "Start Date", "End Date",
])

if "Sheet" in wb.sheetnames:
    del wb["Sheet"]

output_path = "Produksjonsmodell_Mal.xlsx"
wb.save(output_path)
print(f">>> Oppdatert mal lagret: {output_path}")
print(f"    {len(wb.sheetnames)} ark med kolonnebeskrivelser:")
for i, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    antall = sum(1 for c in ws[1] if c.value)
    antall_val = len(ws.data_validations.dataValidation) if ws.data_validations else 0
    print(f"    {i}. {name} ({antall} kolonner, {antall_val} dropdowns)")